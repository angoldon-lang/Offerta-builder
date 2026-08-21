import json
import os
from decimal import Decimal

import pytest

from offerta_builder.docx_builder import extract_text
from offerta_builder.pipeline import BlockingError, build_offer, policy_from_form, prepare_offer
from offerta_builder.qa import LEVEL_FAIL


def test_flusso_completo_senza_template(tmp_path, csv_bom_path, form_data):
    result = build_offer(
        bom_paths=[csv_bom_path], form=form_data, output_dir=str(tmp_path)
    )
    for chiave in ("docx", "bom", "qa", "data", "report"):
        assert os.path.exists(result.outputs[chiave]), chiave
    assert result.qa.status != LEVEL_FAIL
    testo = extract_text(result.outputs["docx"])
    assert "BPER BANCA SPA" in testo
    assert "OFF_ADC_26/0313_R03" in testo
    assert "{{" not in testo


def test_flusso_completo_con_template(tmp_path, csv_bom_path, form_data, template_path):
    result = build_offer(
        bom_paths=[csv_bom_path], form=form_data, output_dir=str(tmp_path),
        template_path=template_path,
    )
    testo = extract_text(result.outputs["docx"])
    assert "2078005" in testo  # le righe della BOM finiscono nella tabella
    assert "{{" not in testo and "{%" not in testo
    assert result.qa.status != LEVEL_FAIL


def test_output_json_coerenti_con_i_totali(tmp_path, csv_bom_path, form_data):
    result = build_offer(
        bom_paths=[csv_bom_path], form=form_data, output_dir=str(tmp_path)
    )
    dati = json.load(open(result.outputs["data"], encoding="utf-8"))
    somma_righe = round(sum(riga["sell_net_total"] for riga in dati["items"]), 2)
    assert somma_righe == dati["totals"]["total_net"]
    qa = json.load(open(result.outputs["qa"], encoding="utf-8"))
    assert qa["status"] in {"ok", "warn", "fail"}


def test_form_incompleto_ferma_il_flusso(tmp_path, csv_bom_path, form_data):
    form_data["condizioni_pagamento"] = ""
    with pytest.raises(BlockingError) as exc:
        build_offer(bom_paths=[csv_bom_path], form=form_data, output_dir=str(tmp_path))
    assert any(i.code == "form.missing_required" for i in exc.value.issues)


def test_bom_incompleta_ferma_il_flusso(tmp_path, form_data):
    bom_path = tmp_path / "bom.csv"
    bom_path.write_text("Codice;Descrizione;Q.ta\nAAA;Licenza;1\n", encoding="utf-8")
    with pytest.raises(BlockingError):
        build_offer(bom_paths=[str(bom_path)], form=form_data, output_dir=str(tmp_path))


def test_force_prosegue_nonostante_le_anomalie(tmp_path, form_data):
    bom_path = tmp_path / "bom.csv"
    bom_path.write_text("Codice;Descrizione;Q.ta\nAAA;Licenza;1\n", encoding="utf-8")
    result = build_offer(
        bom_paths=[str(bom_path)], form=form_data, output_dir=str(tmp_path), force=True
    )
    assert os.path.exists(result.outputs["docx"])
    assert result.qa.status == LEVEL_FAIL  # l'anomalia resta tracciata


def test_pdf_non_generato_per_default(tmp_path, csv_bom_path, form_data):
    """Il deliverable è il DOCX: il PDF si chiede esplicitamente."""
    result = build_offer(bom_paths=[csv_bom_path], form=form_data, output_dir=str(tmp_path))
    assert "pdf" not in result.outputs
    assert not os.path.exists(os.path.join(str(tmp_path), "offerta.pdf"))


def test_pdf_su_richiesta(tmp_path, csv_bom_path, form_data):
    from offerta_builder.pdf import find_soffice

    if not find_soffice():
        pytest.skip("LibreOffice non disponibile")
    result = build_offer(
        bom_paths=[csv_bom_path], form=form_data, output_dir=str(tmp_path), make_pdf=True
    )
    assert os.path.exists(result.outputs["pdf"])


def test_pdf_bloccato_se_il_qa_fallisce(tmp_path, csv_bom_path, form_data):
    """QA rosso: il DOCX esce (va corretto a mano), il PDF no."""
    # validità oltre quella della quotazione distributore (31/07/2026)
    form_data["validita_offerta"] = "15/08/2026"
    result = build_offer(
        bom_paths=[csv_bom_path], form=form_data, output_dir=str(tmp_path), make_pdf=True, force=False,
    )
    assert result.qa.status == LEVEL_FAIL
    assert os.path.exists(result.outputs["docx"])
    assert "pdf" not in result.outputs
    assert any(i.code == "pipeline.pdf_blocked" for i in result.issues)


def test_policy_da_form_e_override_cli(form_data):
    policy = policy_from_form(form_data, {"mode": "target_margin", "target_margin_percent": "45"})
    assert policy.mode == "target_margin"
    assert policy.target_margin_percent == Decimal("45")
    assert policy.vat_percent == Decimal("22")
    assert policy.contract_years == 3


def test_servizi_dal_form_diventano_righe(tmp_path, csv_bom_path, form_data):
    form_data["servizi_aggiuntivi"] = [
        {"descrizione": "Installazione", "quantita": 1, "prezzo_unitario": 2000, "costo_unitario": 800}
    ]
    result = build_offer(
        bom_paths=[csv_bom_path], form=form_data, output_dir=str(tmp_path)
    )
    servizi = [i for i in result.offer.items if i.pricing_mode == "servizio"]
    assert len(servizi) == 1
    assert servizi[0].sell_net_total == Decimal("2000.00")


# ---------------------------------------------------------------------------
# Template Word: celle unite, caselle di testo, assenza di segnaposto
# ---------------------------------------------------------------------------

def _docx_con_cella_unita(percorso):
    from docx import Document

    documento = Document()
    tabella = documento.add_table(rows=2, cols=3)
    unita = tabella.rows[0].cells[0].merge(tabella.rows[0].cells[1])
    unita.text = "TOTALE MATERIALI"
    tabella.rows[1].cells[0].text = "Riga"
    documento.save(percorso)
    return percorso


def test_celle_unite_non_sembrano_titoli_duplicati(tmp_path):
    """python-docx ripete la cella unita per ogni colonna: non è un duplicato."""
    from offerta_builder.docx_builder import extract_text
    from offerta_builder.qa import LEVEL_OK, _check_duplicated_titles

    testo = extract_text(_docx_con_cella_unita(str(tmp_path / "t.docx")))
    assert testo.count("TOTALE MATERIALI") == 1
    assert _check_duplicated_titles(testo).level == LEVEL_OK


def test_template_non_riconosciuto_viene_segnalato(tmp_path, csv_bom_path, form_data):
    """Word senza segnaposto né etichette note: il documento resta vuoto, e si dice."""
    from offerta_builder.docx_builder import has_placeholders

    template = _docx_con_cella_unita(str(tmp_path / "template.docx"))
    assert has_placeholders(template) is False

    result = build_offer(
        bom_paths=[csv_bom_path], form=form_data, output_dir=str(tmp_path / "out"),
        template_path=template, force=True,
    )
    assert any(i.code == "pipeline.template_non_riconosciuto" for i in result.issues)


def test_template_di_esempio_ha_i_segnaposto(template_path):
    from offerta_builder.docx_builder import has_placeholders

    assert has_placeholders(template_path) is True


def test_contesto_separa_prodotti_e_servizi(csv_bom_path, form_data):
    from offerta_builder.bom import normalize
    from offerta_builder.content import build_content
    from offerta_builder.docx_builder import build_context

    form_data["servizi_aggiuntivi"] = [
        {"descrizione": "Installazione", "quantita": 1, "prezzo_unitario": 2000, "costo_unitario": 800}
    ]
    prepared = prepare_offer([normalize(csv_bom_path)], form_data)
    contesto = build_context(prepared.offer, prepared.form, build_content(prepared.offer, prepared.form))

    assert len(contesto["prodotti"]) == 4
    assert len(contesto["servizi"]) == 1
    assert contesto["totale_servizi"] == "2.000,00 EUR"
    assert contesto["prodotti"][0]["descrizione_completa"].startswith("2078005 - SolarWinds")
    assert "31/12/2026" in contesto["prodotti"][0]["descrizione_completa"]


# ---------------------------------------------------------------------------
# Template Word normale (senza segnaposto), compilato per etichette
# ---------------------------------------------------------------------------

def _modello_ad(percorso):
    """Riproduce la struttura del modello AD: tabelle con etichette, niente segnaposto."""
    from docx import Document

    documento = Document()
    documento.add_heading("Oggetto", level=1)
    documento.add_paragraph("[Descrizione]", style="List Bullet")
    documento.add_heading("Requisiti", level=2)
    documento.add_paragraph("…", style="List Bullet")
    documento.add_heading("Esclusioni", level=2)
    documento.add_paragraph("…", style="List Bullet")

    documento.add_heading("Offerta Economica", level=1)
    economica = documento.add_table(rows=8, cols=3)
    economica.style = "Table Grid"
    for indice, testo in enumerate(["Descrizione", "Quantità", "Prezzo €"]):
        economica.rows[0].cells[indice].text = testo
    economica.rows[1].cells[0].text = "[descrizione generica prodotti]"
    economica.rows[2].cells[0].text = "[descrizione di dettaglio]"
    economica.rows[3].cells[0].text = "TOTALE MATERIALI"
    economica.rows[4].cells[0].text = "[descrizione generica servizi]"
    economica.rows[5].cells[0].text = "[descrizione di dettaglio]"
    economica.rows[6].cells[0].text = "TOTALE SERVIZI"
    economica.rows[7].cells[0].text = "Netto a Voi Riservato"

    documento.add_heading("Condizioni di vendita", level=1)
    condizioni = documento.add_table(rows=6, cols=2)
    condizioni.style = "Table Grid"
    for indice, etichetta in enumerate([
        "Tipologia di Pagamento", "Condizioni di Pagamento", "Fatturazione",
        "Validità contratto", "Validità offerta", "Rinnovo",
    ]):
        condizioni.rows[indice].cells[0].text = etichetta
    documento.save(percorso)
    return percorso


def test_word_senza_segnaposto_viene_compilato_per_etichette(tmp_path, csv_bom_path, form_data):
    from offerta_builder.docx_builder import extract_text
    from offerta_builder.template_word import recognized_sections

    template = _modello_ad(str(tmp_path / "modello_ad.docx"))
    assert "offerta economica" in recognized_sections(template)

    form_data["servizi_aggiuntivi"] = [
        {"descrizione": "Installazione", "quantita": 1, "prezzo_unitario": 2000, "costo_unitario": 800}
    ]
    result = build_offer(
        bom_paths=[csv_bom_path], form=form_data, output_dir=str(tmp_path / "out"),
        template_path=template,
    )
    testo = extract_text(result.outputs["docx"])

    assert "2078005" in testo                      # le righe della BOM sono entrate
    assert "Installazione" in testo                # e anche i servizi
    assert "[descrizione di dettaglio]" not in testo
    assert "[descrizione generica prodotti]" not in testo
    assert result.qa.status != "fail"
    assert any(i.code == "pipeline.template_per_etichette" for i in result.issues)


def test_totali_scritti_nelle_righe_giuste(tmp_path, csv_bom_path, form_data):
    from docx import Document

    template = _modello_ad(str(tmp_path / "modello_ad.docx"))
    form_data["servizi_aggiuntivi"] = [
        {"descrizione": "Installazione", "quantita": 1, "prezzo_unitario": 2000, "costo_unitario": 800}
    ]
    result = build_offer(
        bom_paths=[csv_bom_path], form=form_data, output_dir=str(tmp_path / "out"),
        template_path=template,
    )
    tabella = Document(result.outputs["docx"]).tables[0]
    from offerta_builder.money import format_eur

    righe = {r.cells[0].text.strip().lower(): r.cells[-1].text.strip() for r in tabella.rows}
    assert righe["totale servizi"] == "2.000,00 EUR"
    assert righe["netto a voi riservato"] == format_eur(result.offer.totals.total_net)


def test_blocco_servizi_rimosso_se_non_ci_sono_servizi(tmp_path, csv_bom_path, form_data):
    from offerta_builder.docx_builder import extract_text

    template = _modello_ad(str(tmp_path / "modello_ad.docx"))
    form_data["servizi_aggiuntivi"] = []
    result = build_offer(
        bom_paths=[csv_bom_path], form=form_data, output_dir=str(tmp_path / "out"),
        template_path=template,
    )
    testo = extract_text(result.outputs["docx"])
    assert "TOTALE SERVIZI" not in testo
    assert "TOTALE MATERIALI" in testo


def test_condizioni_scritte_accanto_alle_etichette(tmp_path, csv_bom_path, form_data):
    from docx import Document

    template = _modello_ad(str(tmp_path / "modello_ad.docx"))
    result = build_offer(
        bom_paths=[csv_bom_path], form=form_data, output_dir=str(tmp_path / "out"),
        template_path=template,
    )
    condizioni = Document(result.outputs["docx"]).tables[1]
    valori = {r.cells[0].text.strip().lower(): r.cells[1].text.strip() for r in condizioni.rows}
    assert valori["tipologia di pagamento"] == "Bonifico bancario"
    assert valori["condizioni di pagamento"] == "30 gg fine mese"
    assert valori["validità offerta"] == "31/07/2026"
