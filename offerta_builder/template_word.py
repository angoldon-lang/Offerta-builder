"""Compilazione di un template Word **senza segnaposto**.

Il modello AD è un documento normale, scritto a mano in Word: la tabella
dell'offerta economica ha le sue etichette (``[descrizione di dettaglio]``,
``TOTALE MATERIALI``, ``Netto a Voi Riservato``), le condizioni di vendita sono
una tabella etichetta/valore, la copertina sta in caselle di testo.

Questo modulo riconosce quelle etichette e ci scrive dentro i dati dell'offerta,
conservando stili, colori e formattazione del modello. Nessun segnaposto da
inserire: si carica il Word così com'è.

Chi preferisce i segnaposto Jinja continua a usarli: il generatore sceglie da
solo la strada in base al template (vedi ``docx_builder.render``).
"""

from __future__ import annotations

import copy
import re
from typing import Any, Dict, List, Optional

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

# --- copertina: etichetta nel testo -> campo del contesto -------------------
ETICHETTE_COPERTINA = [
    ("offerta per", "cliente"),
    ("spett.le", "cliente"),
    ("p.iva", "piva"),
    ("alla c.a. di", "referente"),
    ("autore:", "autore"),
    ("rif:", "riferimento_offerta"),
    ("la presente offerta è valida fino al", "validita_offerta"),
]
TESTO_DATA_VUOTA = "fare clic o toccare qui per immettere una data"
SOLO_DATA = re.compile(r"^\s*\d{1,2}/\d{1,2}/\d{2,4}\s*$")
# Una data ovunque nella riga: "Modena - 15/10/2026", a volte ripetuta.
DATA_IN_RIGA = re.compile(r"\d{1,2}/\d{1,2}/\d{2,4}")
RIEMPITIVO = re.compile(r"[_\s]{3,}$")

# --- condizioni di vendita: etichetta di riga -> campo ----------------------
ETICHETTE_CONDIZIONI = [
    # Alcuni modelli uniscono tipologia e condizioni in una riga sola.
    ("tipologia e condizioni di pagamento", "pagamento_completo"),
    ("tipologia di pagamento", "tipologia_pagamento"),
    ("condizioni di pagamento", "condizioni_pagamento"),
    ("fatturazione", "fatturazione"),
    ("validità contratto", "durata_contratto"),
    ("validità offerta", "validita_offerta"),
    ("rinnovo", "rinnovo"),
]

# --- tabella economica ------------------------------------------------------
RIGA_DETTAGLIO = "descrizione di dettaglio"
RIGA_GENERICA_PRODOTTI = "descrizione generica prodotti"
RIGA_GENERICA_SERVIZI = "descrizione generica servizi"
RIGA_TOTALE_PRODOTTI = "totale materiali"
RIGA_TOTALE_SERVIZI = "totale servizi"
RIGA_TOTALE_OFFERTA = ("netto a voi riservato", "totale offerta", "totale generale")

SEGNAPOSTO_OGGETTO = ("[descrizione]", "[oggetto]")
PUNTO_ELENCO_VUOTO = ("…", "...")

# Note che il modello rivolge a chi scrive l'offerta: non devono finire al
# cliente, si tolgono sempre.
NOTA_REDAZIONALE = re.compile(r"^\[\s*_*\s*(inserire|indicare|scrivere|nota)\b", re.IGNORECASE)

TITOLI_REQUISITI = ("requisiti",)
TITOLI_ESCLUSIONI = ("esclusioni",)
TITOLO_SEZIONE_OPZIONALE = "requisiti ed esclusioni"


def recognized_sections(template_path: str) -> List[str]:
    """Sezioni che il compilatore sa riconoscere in questo template.

    Serve a dirlo prima di generare: un modello riconosciuto si carica com'è,
    uno sconosciuto va corredato di segnaposto.
    """
    from docx import Document

    try:
        documento = Document(template_path)
    except Exception:  # pragma: no cover - file non leggibile
        return []

    sezioni: List[str] = []
    etichette = [etichetta for etichetta, _ in ETICHETTE_COPERTINA]
    for paragrafo in _paragrafi_caselle(documento):
        testo = paragrafo.text.strip().lower()
        if any(testo.startswith(etichetta) for etichetta in etichette):
            sezioni.append("copertina")
            break

    if _trova_tabella_economica(documento) is not None:
        sezioni.append("offerta economica")

    for tabella in documento.tables:
        for riga in tabella.rows:
            celle = _celle_distinte(riga)
            if len(celle) >= 2 and any(
                celle[0].text.strip().lower().startswith(chiave)
                for chiave, _ in ETICHETTE_CONDIZIONI
            ):
                sezioni.append("condizioni di vendita")
                break
        if "condizioni di vendita" in sezioni:
            break

    for paragrafo in documento.paragraphs:
        testo = paragrafo.text.strip()
        if testo.lower() in SEGNAPOSTO_OGGETTO or testo in PUNTO_ELENCO_VUOTO:
            sezioni.append("oggetto, requisiti ed esclusioni")
            break
    return sezioni


def _aggiorna_sommario(documento) -> None:
    """Chiede a Word di rigenerare i campi (Sommario compreso) all'apertura.

    Togliendo un capitolo la numerazione cambia: senza questo, l'indice a
    pagina 2 resterebbe quello vecchio finché qualcuno non preme F9.
    """
    from docx.oxml.ns import qn

    impostazioni = documento.settings.element
    if impostazioni.find(qn("w:updateFields")) is not None:
        return
    elemento = impostazioni.makeelement(qn("w:updateFields"), {qn("w:val"): "true"})
    impostazioni.append(elemento)


def fill_document(template_path: str, context: Dict[str, Any], output_path: str) -> List[str]:
    """Compila il template riconoscendo le etichette. Ritorna cosa ha riempito."""
    from docx import Document

    documento = Document(template_path)
    fatti: List[str] = []

    if _compila_copertina(documento, context):
        fatti.append("copertina")
    if _ripara_titoli(documento):
        fatti.append("titoli riparati")
    if _compila_tabella_economica(documento, context):
        fatti.append("offerta economica")
    if _compila_condizioni(documento, context):
        fatti.append("condizioni di vendita")
    if _compila_paragrafi(documento, context):
        fatti.append("oggetto, premessa, requisiti ed esclusioni")
    if _elenco_oggetto(documento, context):
        fatti.append("elenco prodotti nell'oggetto")

    _aggiorna_sommario(documento)
    documento.save(output_path)
    return fatti


# ---------------------------------------------------------------------------
# Testo: si scrive nel primo run, così la formattazione del modello resta.
# ---------------------------------------------------------------------------

def _scrivi(paragrafo, testo: str) -> None:
    if paragrafo.runs:
        paragrafo.runs[0].text = testo
        for run in paragrafo.runs[1:]:
            run.text = ""
    else:
        paragrafo.add_run(testo)


def _scrivi_cella(cella, testo: str) -> None:
    _scrivi(cella.paragraphs[0], testo)
    for extra in cella.paragraphs[1:]:
        _scrivi(extra, "")


def _celle_distinte(riga) -> List[Any]:
    viste = set()
    distinte = []
    for cella in riga.cells:
        if id(cella._tc) in viste:
            continue
        viste.add(id(cella._tc))
        distinte.append(cella)
    return distinte


def _testo_riga(riga) -> str:
    return " ".join(cella.text for cella in _celle_distinte(riga)).strip().lower()


def _clona_riga(tabella, modello, valori: List[str], prima: bool = True):
    from docx.table import _Row

    nuovo = copy.deepcopy(modello._tr)
    if prima:
        modello._tr.addprevious(nuovo)
    else:
        modello._tr.addnext(nuovo)
    riga = _Row(nuovo, tabella)
    celle = _celle_distinte(riga)
    for indice, cella in enumerate(celle):
        _scrivi_cella(cella, valori[indice] if indice < len(valori) else "")
    return riga


def _elimina_riga(riga) -> None:
    riga._tr.getparent().remove(riga._tr)


# ---------------------------------------------------------------------------
# Copertina (caselle di testo)
# ---------------------------------------------------------------------------

def _caselle(documento):
    """Caselle di testo del documento, ognuna con i suoi paragrafi in ordine."""
    from docx.text.paragraph import Paragraph

    for casella in documento.element.body.iter(f"{W_NS}txbxContent"):
        yield [Paragraph(elemento, documento) for elemento in casella.iter(f"{W_NS}p")]


def _paragrafi_caselle(documento):
    for paragrafi in _caselle(documento):
        for paragrafo in paragrafi:
            yield paragrafo


def _e_etichetta(testo: str) -> bool:
    spoglio = testo.strip().lower()
    return any(spoglio.startswith(etichetta) for etichetta, _ in ETICHETTE_COPERTINA)


def _compila_copertina(documento, context: Dict[str, Any]) -> bool:
    """Scrive i dati in copertina, sostituendo quelli eventualmente presenti.

    Nel modello AD l'etichetta e il suo valore possono stare su due paragrafi
    diversi ("Spett.le" e sotto la ragione sociale): il valore vecchio va
    sostituito, non affiancato al nuovo.
    """
    scritto = False
    for paragrafi in _caselle(documento):
        pieni = [p for p in paragrafi if p.text.strip()]
        for posizione, paragrafo in enumerate(pieni):
            testo = paragrafo.text.strip()
            spoglio = testo.lower()

            if spoglio.startswith(TESTO_DATA_VUOTA):
                _scrivi(paragrafo, str(context.get("validita_offerta", "")))
                scritto = True
                continue

            # "Modena - 15/10/2026" (a volte con la data ripetuta due volte)
            nuova_data = str(context.get("data_offerta", "")).strip()
            if nuova_data and DATA_IN_RIGA.search(testo):
                ripulito = DATA_IN_RIGA.sub("", testo).strip()
                if not _e_etichetta(testo) or spoglio.startswith("modena"):
                    _scrivi(paragrafo, f"{ripulito} {nuova_data}".strip())
                    scritto = True
                    continue

            for etichetta, campo in ETICHETTE_COPERTINA:
                valore = str(context.get(campo, "")).strip()
                if not valore or not spoglio.startswith(etichetta):
                    continue
                resto = testo[len(etichetta):].strip(" :_\t")
                if resto:
                    # etichetta e valore sulla stessa riga: si riscrive il valore
                    _scrivi(paragrafo, f"{testo[:len(etichetta)].strip()} {valore}")
                elif posizione + 1 < len(pieni) and not _e_etichetta(pieni[posizione + 1].text):
                    # il valore sta nel paragrafo sotto
                    _scrivi(pieni[posizione + 1], valore)
                else:
                    _scrivi(paragrafo, f"{testo.strip()} {valore}")
                scritto = True
                break
    return scritto


def _ripara_titoli(documento) -> bool:
    """Ricompone i titoli finiti doppi ("PremessaPremessa").

    Capita quando un documento viene compilato aggiungendo un run invece di
    riscrivere quello che c'è: il titolo risulta scritto due volte di fila.
    """
    riparati = False
    paragrafi = list(documento.paragraphs)
    for indice, paragrafo in enumerate(paragrafi):
        stile = (paragrafo.style.name if paragrafo.style is not None else "").lower()
        if not stile.startswith(("heading", "titolo")):
            continue

        testi = [run.text.strip() for run in paragrafo.runs if run.text.strip()]
        if len(testi) >= 2 and all(t.lower() == testi[0].lower() for t in testi[1:]):
            _scrivi(paragrafo, testi[0])
            riparati = True

        # Titolo ripetuto come testo normale nella riga sotto.
        titolo = paragrafo.text.strip().lower()
        successivo = paragrafi[indice + 1] if indice + 1 < len(paragrafi) else None
        if successivo is None or not titolo:
            continue
        stile_succ = (successivo.style.name if successivo.style is not None else "").lower()
        if stile_succ.startswith(("heading", "titolo")):
            continue
        if successivo.text.strip().lower() == titolo:
            _elimina_paragrafo(successivo)
            riparati = True
    return riparati


# ---------------------------------------------------------------------------
# Tabella dell'offerta economica
# ---------------------------------------------------------------------------

def _indice_riga(tabella, *etichette: str) -> Optional[int]:
    for indice, riga in enumerate(tabella.rows):
        testo = _testo_riga(riga)
        if any(etichetta in testo for etichetta in etichette):
            return indice
    return None


def _trova_tabella_economica(documento):
    for tabella in documento.tables:
        testo = " ".join(_testo_riga(riga) for riga in tabella.rows)
        if RIGA_DETTAGLIO in testo or RIGA_TOTALE_OFFERTA[0] in testo:
            return tabella
    return None


def _righe_gia_compilate(tabella) -> Dict[str, List[int]]:
    """Righe dati di una tabella già compilata, divise fra materiali e servizi.

    Serve per riallineare una vecchia offerta: le sue righe vengono sostituite
    da quelle nuove, conservando la formattazione della prima.
    """
    gruppi: Dict[str, List[int]] = {"prodotti": [], "servizi": []}
    corrente = "prodotti"
    for indice, riga in enumerate(tabella.rows):
        if indice == 0:
            continue
        testo = _testo_riga(riga)
        if any(testo.startswith(etichetta) for etichetta in RIGA_TOTALE_OFFERTA):
            break
        if testo.startswith(RIGA_TOTALE_PRODOTTI):
            corrente = "servizi"
            continue
        if testo.startswith(RIGA_TOTALE_SERVIZI):
            corrente = ""
            continue
        celle = _celle_distinte(riga)
        if not corrente or not celle or not celle[0].text.strip():
            continue
        if any(cella.text.strip() for cella in celle[1:]):
            gruppi[corrente].append(indice)
    return gruppi


def _sostituisci_blocco(tabella, indici: List[int], righe: List[Dict[str, Any]]) -> None:
    """Rimpiazza le righe esistenti di un blocco con quelle nuove.

    Le righe da togliere si prendono come oggetti prima di inserire le nuove:
    gli indici slitterebbero a ogni inserimento.
    """
    da_togliere = [tabella.rows[indice] for indice in indici]
    modello = da_togliere[0]
    for riga in righe:
        _clona_riga(tabella, modello, [
            riga.get("descrizione_completa") or riga.get("descrizione", ""),
            riga.get("quantita", ""),
            riga.get("totale", ""),
        ])
    for riga_vecchia in da_togliere:
        _elimina_riga(riga_vecchia)


def _scrivi_totali(tabella, context: Dict[str, Any]) -> None:
    for etichetta, valore in (
        (RIGA_TOTALE_PRODOTTI, context.get("totale_prodotti", "")),
        (RIGA_TOTALE_SERVIZI, context.get("totale_servizi", "")),
    ):
        indice = _indice_riga(tabella, etichetta)
        if indice is not None:
            _scrivi_cella(_celle_distinte(tabella.rows[indice])[-1], str(valore))
    indice_offerta = _indice_riga(tabella, *RIGA_TOTALE_OFFERTA)
    if indice_offerta is not None:
        _scrivi_cella(
            _celle_distinte(tabella.rows[indice_offerta])[-1],
            str(context.get("totale_imponibile", "")),
        )


def _compila_tabella_economica(documento, context: Dict[str, Any]) -> bool:
    tabella = _trova_tabella_economica(documento)
    if tabella is None:
        return False

    prodotti = context.get("prodotti") or []
    servizi = context.get("servizi") or []

    # I due blocchi hanno la stessa forma: riga generica, righe di dettaglio,
    # totale. Si compilano dal fondo, così gli indici a monte restano validi.
    blocchi = [
        (RIGA_GENERICA_SERVIZI, RIGA_TOTALE_SERVIZI, servizi, context.get("nota_servizi", ""),
         context.get("totale_servizi", "")),
        (RIGA_GENERICA_PRODOTTI, RIGA_TOTALE_PRODOTTI, prodotti, context.get("oggetto", ""),
         context.get("totale_prodotti", "")),
    ]

    dettagli = [
        indice for indice, riga in enumerate(tabella.rows)
        if RIGA_DETTAGLIO in _testo_riga(riga)
    ]

    if not dettagli:
        # Tabella già compilata: è il caso del rinnovo, dove la vecchia offerta
        # fa da modello. Si sostituiscono le righe che ci sono già.
        gruppi = _righe_gia_compilate(tabella)
        if not gruppi["prodotti"] and not gruppi["servizi"]:
            return False
        for nome, righe_nuove in (("servizi", servizi), ("prodotti", prodotti)):
            indici = gruppi[nome]
            if not indici:
                continue
            if righe_nuove:
                _sostituisci_blocco(tabella, indici, righe_nuove)
            else:
                for riga_vecchia in [tabella.rows[indice] for indice in indici]:
                    _elimina_riga(riga_vecchia)
        _scrivi_totali(tabella, context)
        return True

    for posizione, (etichetta_generica, etichetta_totale, righe, testo_generico, totale) in enumerate(blocchi):
        modello_indice = dettagli[-1 - posizione] if len(dettagli) > posizione else None
        if modello_indice is None:
            continue
        modello = tabella.rows[modello_indice]

        if righe:
            for riga in righe:
                _clona_riga(tabella, modello, [
                    riga.get("descrizione_completa") or riga.get("descrizione", ""),
                    riga.get("quantita", ""),
                    riga.get("totale", ""),
                ])
        _elimina_riga(modello)

        indice_generica = _indice_riga(tabella, etichetta_generica)
        if indice_generica is not None:
            if righe:
                _scrivi_cella(_celle_distinte(tabella.rows[indice_generica])[0], str(testo_generico))
            else:
                _elimina_riga(tabella.rows[indice_generica])

        indice_totale = _indice_riga(tabella, etichetta_totale)
        if indice_totale is not None:
            if righe:
                _scrivi_cella(_celle_distinte(tabella.rows[indice_totale])[-1], str(totale))
            else:
                _elimina_riga(tabella.rows[indice_totale])

    # Con un blocco solo il subtotale ripete il totale: si toglie.
    if prodotti and not servizi:
        indice_subtotale = _indice_riga(tabella, RIGA_TOTALE_PRODOTTI)
        if indice_subtotale is not None:
            _elimina_riga(tabella.rows[indice_subtotale])

    indice_offerta = _indice_riga(tabella, *RIGA_TOTALE_OFFERTA)
    if indice_offerta is not None:
        _scrivi_cella(
            _celle_distinte(tabella.rows[indice_offerta])[-1],
            str(context.get("totale_imponibile", "")),
        )
    return True


# ---------------------------------------------------------------------------
# Condizioni di vendita
# ---------------------------------------------------------------------------

def _norm_etichetta(testo: str) -> str:
    spoglio = str(testo).strip().lower()
    for accentata, piana in (("à", "a"), ("è", "e"), ("é", "e"), ("ì", "i"), ("ò", "o"), ("ù", "u")):
        spoglio = spoglio.replace(accentata, piana)
    return spoglio


def _e_tabella_condizioni(tabella) -> bool:
    """Riconosce la tabella etichetta/valore delle condizioni di vendita.

    Senza questo controllo una riga dell'offerta economica che comincia per
    "RINNOVO ..." verrebbe scambiata per la riga "Rinnovo" delle condizioni, e
    il valore finirebbe nella colonna della quantità.
    """
    righe_due_colonne = 0
    etichette = 0
    for riga in tabella.rows:
        celle = _celle_distinte(riga)
        if len(celle) != 2:
            continue
        righe_due_colonne += 1
        testo = _norm_etichetta(celle[0].text)
        if any(testo.startswith(_norm_etichetta(chiave)) for chiave, _ in ETICHETTE_CONDIZIONI):
            etichette += 1
    return righe_due_colonne >= 2 and etichette >= 2


def _compila_condizioni(documento, context: Dict[str, Any]) -> bool:
    """Compila la tabella delle condizioni di vendita.

    Le righe previste dal modello ma senza valore (tipicamente "Validità
    contratto" e "Rinnovo" in una fornitura di licenze) vengono tolte: meglio
    una tabella corta che una con celle vuote.
    """
    scritto = False
    for tabella in documento.tables:
        if not _e_tabella_condizioni(tabella):
            continue
        da_togliere = []
        for riga in tabella.rows:
            celle = _celle_distinte(riga)
            if len(celle) < 2:
                continue
            etichetta = _norm_etichetta(celle[0].text)
            if any(etichetta.startswith(_norm_etichetta(x)) for x in context.get("condizioni_da_togliere", [])):
                da_togliere.append(riga)
                continue
            for chiave, campo in ETICHETTE_CONDIZIONI:
                if not etichetta.startswith(_norm_etichetta(chiave)):
                    continue
                valore = str(context.get(campo, "")).strip()
                if valore:
                    if campo == "validita_offerta" and "fino al" in celle[1].text.strip().lower():
                        valore = f"FINO AL {valore}"
                    _scrivi_cella(celle[1], valore)
                    scritto = True
                else:
                    da_togliere.append(riga)
                break
        for riga in da_togliere:
            _elimina_riga(riga)
            scritto = True
    return scritto


# ---------------------------------------------------------------------------
# Oggetto, premessa, requisiti, esclusioni
# ---------------------------------------------------------------------------

def _sezione_corrente(paragrafi, indice: int) -> str:
    for precedente in reversed(paragrafi[:indice]):
        stile = (precedente.style.name if precedente.style is not None else "").lower()
        if stile.startswith("heading") and precedente.text.strip():
            return precedente.text.strip().lower()
    return ""


def _espandi_elenco(paragrafo, voci: List[str]) -> None:
    """Sostituisce un punto elenco segnaposto con una voce per riga."""
    from docx.text.paragraph import Paragraph

    if not voci:
        _scrivi(paragrafo, "")
        return
    _scrivi(paragrafo, str(voci[0]))
    riferimento = paragrafo
    for voce in voci[1:]:
        nuovo = copy.deepcopy(paragrafo._p)
        riferimento._p.addnext(nuovo)
        riferimento = Paragraph(nuovo, paragrafo._parent)
        _scrivi(riferimento, str(voce))


def _livello_titolo(paragrafo) -> int:
    stile = (paragrafo.style.name if paragrafo.style is not None else "").lower()
    match = re.search(r"(?:heading|titolo)\s*(\d)", stile)
    return int(match.group(1)) if match else 0


def _elimina_paragrafo(paragrafo) -> None:
    elemento = paragrafo._p
    genitore = elemento.getparent()
    if genitore is not None:
        genitore.remove(elemento)


def _elimina_sezione(paragrafi, indice: int) -> None:
    """Toglie un titolo e tutto quello che sta sotto, fino al titolo pari o superiore."""
    livello = _livello_titolo(paragrafi[indice])
    _elimina_paragrafo(paragrafi[indice])
    for successivo in paragrafi[indice + 1:]:
        prossimo_livello = _livello_titolo(successivo)
        if prossimo_livello and prossimo_livello <= livello:
            break
        _elimina_paragrafo(successivo)


def _rimuovi_sezioni_vuote(documento, context: Dict[str, Any]) -> None:
    """Toglie Requisiti ed Esclusioni quando non ci sono voci da elencare.

    Il modello AD le prevede "solo se necessario": lasciarle vuote significa non
    volerle, e restare con un titolo senza contenuto fa brutta figura in offerta.
    """
    requisiti = context.get("requisiti") or []
    esclusioni = context.get("esclusioni") or []
    if requisiti and esclusioni:
        return

    paragrafi = list(documento.paragraphs)
    for indice, paragrafo in enumerate(paragrafi):
        titolo = paragrafo.text.strip().lower()
        if not _livello_titolo(paragrafo) or not titolo:
            continue
        if titolo.startswith(TITOLO_SEZIONE_OPZIONALE) and not requisiti and not esclusioni:
            _elimina_sezione(paragrafi, indice)
            return
        if titolo.startswith(TITOLI_REQUISITI) and not requisiti and titolo != TITOLO_SEZIONE_OPZIONALE:
            _elimina_sezione(paragrafi, indice)
        elif titolo.startswith(TITOLI_ESCLUSIONI) and not esclusioni:
            _elimina_sezione(paragrafi, indice)


def _e_punto_elenco(paragrafo) -> bool:
    """Un paragrafo è un punto elenco se ha la numerazione o uno stile da elenco.

    Il pallino può arrivare dal paragrafo (``w:numPr``, come nei modelli AD) o
    dallo stile ("List Bullet"): vanno riconosciuti entrambi.
    """
    from docx.oxml.ns import qn

    pPr = paragrafo._p.pPr
    if pPr is not None and pPr.find(qn("w:numPr")) is not None:
        return True
    stile = (paragrafo.style.name if paragrafo.style is not None else "").lower()
    return stile.startswith(("list ", "elenco"))


def _donatore_elenco(documento):
    """Un punto elenco già presente nel modello, da cui copiare bullet e rientri.

    Lo stile "List Paragraph" da solo non disegna il pallino: nei modelli AD la
    numerazione sta sul paragrafo (``w:numPr``). Copiare un punto elenco vero è
    l'unico modo di ottenere un elenco identico agli altri del documento.
    """
    ripiego = None
    for paragrafo in documento.paragraphs:
        if _livello_titolo(paragrafo) or not paragrafo.text.strip():
            continue
        if not _e_punto_elenco(paragrafo):
            continue
        from docx.oxml.ns import qn

        pPr = paragrafo._p.pPr
        if pPr is not None and pPr.find(qn("w:numPr")) is not None:
            return paragrafo
        ripiego = ripiego or paragrafo
    return ripiego


def _elenco_oggetto(documento, context: Dict[str, Any]) -> bool:
    """Elenca i prodotti sotto il capitolo Oggetto, anche senza segnaposto.

    Il modello AD scrive "...la fornitura dei seguenti prodotti e servizi:" e
    poi lascia lo spazio bianco: l'elenco va aggiunto, altrimenti il capitolo
    annuncia un elenco che non c'è.
    """
    from docx.text.paragraph import Paragraph
    import copy as _copy

    voci = [str(v).strip() for v in (context.get("voci_oggetto") or []) if str(v).strip()]
    if not voci:
        return False

    paragrafi = list(documento.paragraphs)
    for indice, paragrafo in enumerate(paragrafi):
        if not _livello_titolo(paragrafo) or paragrafo.text.strip().lower() != "oggetto":
            continue

        sezione = []
        for successivo in paragrafi[indice + 1:]:
            if _livello_titolo(successivo):
                break
            sezione.append(successivo)
        pieni = [par for par in sezione if par.text.strip()]
        if not pieni:
            return False
        # Già elencato (segnaposto espanso o documento ricompilato): non si tocca.
        gia_scritte = {par.text.strip() for par in pieni}
        if any(voce in gia_scritte for voce in voci):
            return False

        donatore = _donatore_elenco(documento)
        modello = donatore._p if donatore is not None else pieni[-1]._p
        riferimento = pieni[-1]
        for voce in voci:
            nuovo_p = _copy.deepcopy(modello)
            riferimento._p.addnext(nuovo_p)
            riferimento = Paragraph(nuovo_p, paragrafo._parent)
            _scrivi(riferimento, voce)
        return True
    return False


def _compila_paragrafi(documento, context: Dict[str, Any]) -> bool:
    paragrafi = list(documento.paragraphs)
    scritto = False

    for paragrafo in paragrafi:
        if NOTA_REDAZIONALE.match(paragrafo.text.strip()):
            _elimina_paragrafo(paragrafo)
            scritto = True

    paragrafi = list(documento.paragraphs)
    for indice, paragrafo in enumerate(paragrafi):
        testo = paragrafo.text.strip()
        spoglio = testo.lower()
        stile = (paragrafo.style.name if paragrafo.style is not None else "").lower()

        if spoglio in SEGNAPOSTO_OGGETTO:
            _scrivi(paragrafo, str(context.get("oggetto", "")))
            scritto = True
        elif testo in PUNTO_ELENCO_VUOTO:
            sezione = _sezione_corrente(paragrafi, indice)
            if sezione.startswith("requisiti") and "esclusioni" not in sezione:
                _espandi_elenco(paragrafo, context.get("requisiti") or [])
            elif sezione.startswith("esclusioni"):
                _espandi_elenco(paragrafo, context.get("esclusioni") or [])
            elif context.get("voci_oggetto"):
                _espandi_elenco(paragrafo, context.get("voci_oggetto") or [])
            else:
                _scrivi(paragrafo, str(context.get("descrizione_fornitura", "")))
            scritto = True
        elif stile.startswith("heading") and spoglio == "premessa":
            successivo = paragrafi[indice + 1] if indice + 1 < len(paragrafi) else None
            if successivo is not None and not successivo.text.strip():
                _scrivi(successivo, str(context.get("premessa", "")))
                scritto = True

    _rimuovi_sezioni_vuote(documento, context)
    return scritto
