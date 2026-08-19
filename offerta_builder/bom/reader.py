"""Lettura dei file BOM: CSV/TSV, XLSX e PDF.

Ogni lettore restituisce delle ``RawTable``; la scelta della tabella giusta e
la mappatura sullo schema unico avvengono nel normalizzatore.
"""

from __future__ import annotations

import csv
import os
import re
from typing import Dict, List, Tuple

from .base import RawTable, clean_cell, is_empty_row
from .columns import map_header

SUPPORTED_EXTENSIONS = {".csv", ".tsv", ".txt", ".xlsx", ".xlsm", ".pdf"}


def read_tables(path: str) -> List[RawTable]:
    """Estrae tutte le tabelle candidate da un file BOM."""
    ext = os.path.splitext(path)[1].lower()
    if ext in {".csv", ".tsv", ".txt"}:
        return _read_csv(path)
    if ext in {".xlsx", ".xlsm"}:
        return _read_xlsx(path)
    if ext == ".pdf":
        return _read_pdf(path)
    raise ValueError(
        f"Formato non supportato: {ext or path}. "
        f"Formati ammessi: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
    )


def _split_matrix(matrix: List[List[str]], source: str, fmt: str, sheet: str = "") -> List[RawTable]:
    """Individua la riga di intestazione e taglia il corpo tabella."""
    tables: List[RawTable] = []
    best: Tuple[int, Dict[int, str], int] = (-1, {}, 0)
    for idx, row in enumerate(matrix[:40]):
        mapping, score = map_header([clean_cell(c) for c in row])
        if score > best[2]:
            best = (idx, mapping, score)
    header_idx, mapping, score = best
    if header_idx < 0 or score <= 0:
        return tables

    header_cells = [clean_cell(c) for c in matrix[header_idx]]
    header = [mapping.get(i, header_cells[i] if i < len(header_cells) else "") for i in range(len(header_cells))]
    body: List[List[str]] = []
    for row in matrix[header_idx + 1:]:
        cells = [clean_cell(c) for c in row]
        if is_empty_row(cells):
            continue
        body.append(cells)

    meta_text = "\n".join(
        " ".join(clean_cell(c) for c in row) for row in matrix[:header_idx]
    )
    tables.append(
        RawTable(
            source=source,
            fmt=fmt,
            header=header,
            rows=body,
            meta_text=meta_text,
            sheet=sheet,
            score=score,
        )
    )
    return tables


def _read_csv(path: str) -> List[RawTable]:
    with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        sample = handle.read(8192)
        handle.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ";" if sample.count(";") > sample.count(",") else ","
        matrix = [row for row in csv.reader(handle, delimiter=delimiter)]
    tables = _split_matrix(matrix, path, "csv")
    if not tables:
        # Nessuna intestazione riconosciuta: restituiamo comunque il grezzo,
        # così il normalizzatore può segnalare il problema all'utente.
        tables.append(RawTable(source=path, fmt="csv", header=[], rows=matrix, score=0))
    return tables


def _read_xlsx(path: str) -> List[RawTable]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=True)
    tables: List[RawTable] = []
    try:
        for sheet in workbook.worksheets:
            matrix: List[List[str]] = []
            for row in sheet.iter_rows(values_only=True):
                matrix.append([clean_cell(cell) for cell in row])
            if not matrix:
                continue
            tables.extend(_split_matrix(matrix, path, "xlsx", sheet=sheet.title))
    finally:
        workbook.close()
    return tables


_PDF_META_STOP = re.compile(r"^\s*$")


def _read_pdf(path: str) -> List[RawTable]:
    import pdfplumber

    tables: List[RawTable] = []
    text_chunks: List[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            text_chunks.append(page_text)
            for raw in page.extract_tables():
                matrix = [[clean_cell(cell) for cell in row] for row in raw if row]
                if len(matrix) < 2:
                    continue
                tables.extend(_split_matrix(matrix, path, "pdf"))
    full_text = "\n".join(text_chunks)

    if not tables:
        tables.extend(_tables_from_text(full_text, path))

    for table in tables:
        # La testata del PDF (numero quote, end user, validità) sta nel testo.
        table.meta_text = (table.meta_text + "\n" + full_text).strip()
    return tables


def _tables_from_text(text: str, path: str) -> List[RawTable]:
    """Fallback per PDF senza tabelle vettoriali: righe allineate a spazi."""
    lines = [line for line in text.splitlines() if line.strip()]
    matrix = [re.split(r"\s{2,}", line.strip()) for line in lines]
    matrix = [row for row in matrix if len(row) >= 3]
    if not matrix:
        return []
    tables = _split_matrix(matrix, path, "pdf")
    for table in tables:
        table.meta_text = text
    return tables
