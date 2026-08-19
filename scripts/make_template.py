"""Genera un template Word di esempio con i segnaposto attesi dal builder.

Serve da riferimento per adattare il template AD reale: i nomi dei segnaposto
sono quelli prodotti da ``docx_builder.build_context``.
"""

from __future__ import annotations

import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

OUTPUT = os.path.join("templates", "offerta_ad_template.docx")


def build(path: str = OUTPUT) -> str:
    document = Document()

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("OFFERTA COMMERCIALE")
    run.bold = True
    run.font.size = Pt(22)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("{{ oggetto }}").bold = True

    cover = document.add_table(rows=0, cols=2)
    cover.style = "Table Grid"
    for label, placeholder in [
        ("Cliente", "{{ cliente }}"),
        ("P.IVA", "{{ piva }}"),
        ("Indirizzo", "{{ indirizzo_cliente }}"),
        ("Referente", "{{ referente }}"),
        ("Riferimento offerta", "{{ riferimento_offerta }}"),
        ("Data offerta", "{{ data_offerta }}"),
        ("Validità offerta", "{{ validita_offerta }}"),
        ("Autore", "{{ autore }}"),
    ]:
        cells = cover.add_row().cells
        cells[0].paragraphs[0].add_run(label).bold = True
        cells[1].text = placeholder

    document.add_page_break()

    document.add_heading("Oggetto", level=1)
    document.add_paragraph("{{ oggetto }}")

    document.add_heading("Premessa", level=1)
    document.add_paragraph("{{ premessa }}")

    document.add_heading("Descrizione della fornitura", level=1)
    document.add_paragraph("{{ descrizione_fornitura }}")

    document.add_heading("Servizi professionali", level=1)
    document.add_paragraph("{%p for servizio in descrizione_servizi %}")
    document.add_paragraph("{{ servizio }}", style="List Bullet")
    document.add_paragraph("{%p endfor %}")

    document.add_heading("Offerta economica", level=1)
    headers = ["#", "Codice", "Descrizione", "Periodo", "Q.ta", "Prezzo unitario", "Totale"]
    # Le righe che contengono {%tr ... %} vengono sostituite dal tag: i marcatori
    # del ciclo stanno quindi in righe dedicate, la riga dati sta in mezzo.
    economics = document.add_table(rows=4, cols=len(headers))
    economics.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = economics.rows[0].cells[index]
        cell.text = ""
        cell.paragraphs[0].add_run(header).bold = True
    economics.rows[1].cells[0].text = "{%tr for riga in righe %}"
    for index, placeholder in enumerate(
        [
            "{{ riga.n }}", "{{ riga.codice }}", "{{ riga.descrizione }}", "{{ riga.periodo }}",
            "{{ riga.quantita }}", "{{ riga.prezzo_unitario }}", "{{ riga.totale }}",
        ]
    ):
        economics.rows[2].cells[index].text = placeholder
    economics.rows[3].cells[0].text = "{%tr endfor %}"

    document.add_paragraph()
    totals = document.add_table(rows=3, cols=2)
    totals.style = "Table Grid"
    for row_index, (label, placeholder) in enumerate(
        [
            ("Totale imponibile", "{{ totale_imponibile }}"),
            ("IVA {{ iva_percento }}", "{{ totale_iva }}"),
            ("Totale offerta", "{{ totale_offerta }}"),
        ]
    ):
        cells = totals.rows[row_index].cells
        cells[0].paragraphs[0].add_run(label).bold = True
        cells[1].paragraphs[0].add_run(placeholder).bold = True

    document.add_heading("Riepilogo per annualità", level=2)
    annual = document.add_table(rows=4, cols=4)
    annual.style = "Table Grid"
    for index, header in enumerate(["Periodo", "Imponibile", "IVA", "Totale"]):
        cell = annual.rows[0].cells[index]
        cell.text = ""
        cell.paragraphs[0].add_run(header).bold = True
    annual.rows[1].cells[0].text = "{%tr for anno in annualita %}"
    for index, placeholder in enumerate(
        ["{{ anno.periodo }}", "{{ anno.imponibile }}", "{{ anno.iva }}", "{{ anno.totale }}"]
    ):
        annual.rows[2].cells[index].text = placeholder
    annual.rows[3].cells[0].text = "{%tr endfor %}"

    document.add_heading("Condizioni di vendita", level=1)
    conditions = document.add_table(rows=0, cols=2)
    conditions.style = "Table Grid"
    for label, placeholder in [
        ("Validità offerta", "{{ validita_offerta }}"),
        ("Tipologia di pagamento", "{{ tipologia_pagamento }}"),
        ("Condizioni di pagamento", "{{ condizioni_pagamento }}"),
        ("Fatturazione", "{{ fatturazione }}"),
        ("Durata contratto", "{{ durata_contratto }}"),
        ("Rinnovo", "{{ rinnovo }}"),
    ]:
        cells = conditions.add_row().cells
        cells[0].paragraphs[0].add_run(label).bold = True
        cells[1].text = placeholder

    document.add_heading("Requisiti a carico del Cliente", level=2)
    document.add_paragraph("{%p for requisito in requisiti %}")
    document.add_paragraph("{{ requisito }}", style="List Bullet")
    document.add_paragraph("{%p endfor %}")

    document.add_heading("Esclusioni", level=2)
    document.add_paragraph("{%p for esclusione in esclusioni %}")
    document.add_paragraph("{{ esclusione }}", style="List Bullet")
    document.add_paragraph("{%p endfor %}")

    document.add_heading("Note commerciali", level=2)
    document.add_paragraph("{%p for nota in note_commerciali %}")
    document.add_paragraph("{{ nota }}", style="List Bullet")
    document.add_paragraph("{%p endfor %}")

    document.add_heading("Accettazione", level=1)
    document.add_paragraph(
        "Per accettazione della presente offerta, comprensiva delle condizioni di vendita e "
        "degli allegati richiamati, timbro e firma del Cliente:"
    )
    signature = document.add_table(rows=2, cols=2)
    signature.style = "Table Grid"
    signature.rows[0].cells[0].text = "Data"
    signature.rows[0].cells[1].text = "Timbro e firma"

    document.add_heading("Allegati", level=1)
    document.add_paragraph("{%p for allegato in allegati %}")
    document.add_paragraph("{{ allegato }}", style="List Bullet")
    document.add_paragraph("{%p endfor %}")

    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    document.save(path)
    return path


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else OUTPUT
    print(build(target))
