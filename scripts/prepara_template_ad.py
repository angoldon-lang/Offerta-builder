"""Inserisce i segnaposto in un template Word AD che ne è privo.

Un template senza segnaposto viene copiato tale e quale: il documento esce
identico al modello, senza i dati dell'offerta. Questo script prende il modello
AD (copertina in caselle di testo, tabella "Offerta Economica" a tre colonne,
tabella "Condizioni di vendita") e ci inserisce i campi che il generatore sa
compilare.

Uso::

    python scripts/prepara_template_ad.py modello_ad.docx templates/offerta_ad.docx

Va rilanciato solo se il modello cambia struttura: il file prodotto è un
template normale, da caricare nell'interfaccia.
"""

from __future__ import annotations

import copy
import os
import re
import sys
import zipfile
from typing import Dict, List

# Copertina: testi presenti nelle caselle di testo -> testo con i segnaposto.
COPERTINA: Dict[str, str] = {
    "Offerta per": "Offerta per {{ cliente }}",
    "Spett.le": "Spett.le {{ cliente }}",
    "P.IVA": "P.IVA {{ piva }}",
    "Alla C.A. di": "Alla C.A. di {{ referente }}",
    "Autore:": "Autore: {{ autore }}",
    "Rif:": "Rif: {{ riferimento_offerta }}",
    "La presente offerta è valida fino al": "La presente offerta è valida fino al {{ validita_offerta }}",
    "Fare clic o toccare qui per immettere una data.": "{{ validita_offerta }}",
}

# Condizioni di vendita: etichetta di riga -> segnaposto da mettere a fianco.
CONDIZIONI: Dict[str, str] = {
    "tipologia di pagamento": "{{ tipologia_pagamento }}",
    "condizioni di pagamento": "{{ condizioni_pagamento }}",
    "fatturazione": "{{ fatturazione }}",
    "validità contratto": "{{ durata_contratto }}",
    "validità offerta": "{{ validita_offerta }}",
    "rinnovo": "{{ rinnovo }}",
}

DATA_LETTERALE = re.compile(r"(Modena\s*[–-]\s*)\d{1,2}/\d{1,2}/\d{2,4}")
# In copertina la data di esempio sta in un run tutto suo: si sostituisce solo
# la prima che si incontra, per non toccare le date citate nelle condizioni.
SOLO_DATA = re.compile(r"^\s*\d{1,2}/\d{1,2}/\d{2,4}\s*$")


# ---------------------------------------------------------------------------
# Copertina (caselle di testo): si lavora sull'XML, python-docx non le espone.
# ---------------------------------------------------------------------------

def _patch_copertina(sorgente: str, destinazione: str) -> None:
    with zipfile.ZipFile(sorgente) as ingresso:
        parti = {nome: ingresso.read(nome) for nome in ingresso.namelist()}
        info = {i.filename: i for i in ingresso.infolist()}

    xml = parti["word/document.xml"].decode("utf-8")

    data_gia_sostituita = {"fatto": False}

    def sostituisci(match: re.Match) -> str:
        apertura, testo, chiusura = match.group(1), match.group(2), match.group(3)
        nuovo = testo
        if SOLO_DATA.match(testo) and not data_gia_sostituita["fatto"]:
            data_gia_sostituita["fatto"] = True
            nuovo = "{{ data_offerta }}"
        for etichetta, rimpiazzo in COPERTINA.items():
            if etichetta in nuovo and "{{" not in nuovo:
                nuovo = nuovo.replace(etichetta, rimpiazzo)
        nuovo = DATA_LETTERALE.sub(r"\1{{ data_offerta }}", nuovo)
        if nuovo == testo:
            return match.group(0)
        if "xml:space" not in apertura:
            apertura = apertura[:-1] + ' xml:space="preserve">'
        return f"{apertura}{nuovo}{chiusura}"

    xml = re.sub(r"(<w:t[^>]*>)([^<]*)(</w:t>)", sostituisci, xml)
    parti["word/document.xml"] = xml.encode("utf-8")

    with zipfile.ZipFile(destinazione, "w", zipfile.ZIP_DEFLATED) as uscita:
        for nome, contenuto in parti.items():
            uscita.writestr(info[nome], contenuto)


# ---------------------------------------------------------------------------
# Corpo del documento
# ---------------------------------------------------------------------------

def _scrivi(paragrafo, testo: str) -> None:
    """Sostituisce il testo del paragrafo conservandone la formattazione."""
    if paragrafo.runs:
        paragrafo.runs[0].text = testo
        for run in paragrafo.runs[1:]:
            run.text = ""
    else:
        paragrafo.add_run(testo)


def _scrivi_cella(cella, testo: str) -> None:
    _scrivi(cella.paragraphs[0], testo)
    for extra in cella.paragraphs[1:]:
        _scrivi(extra, "")


def _celle_distinte(riga):
    """Celle della riga senza ripetere quelle unite."""
    viste = set()
    distinte = []
    for cella in riga.cells:
        if id(cella._tc) in viste:
            continue
        viste.add(id(cella._tc))
        distinte.append(cella)
    return distinte


def _scrivi_riga(riga, testi: List[str]) -> None:
    for indice, cella in enumerate(_celle_distinte(riga)):
        _scrivi_cella(cella, testi[indice] if indice < len(testi) else "")


def _clona(tabella, modello, testi: List[str], prima: bool):
    from docx.table import _Row

    nuovo_tr = copy.deepcopy(modello._tr)
    if prima:
        modello._tr.addprevious(nuovo_tr)
    else:
        modello._tr.addnext(nuovo_tr)
    riga = _Row(nuovo_tr, tabella)
    _scrivi_riga(riga, testi)
    return riga


def _trova_tabella(documento, intestazioni: List[str]):
    attese = [t.lower() for t in intestazioni]
    for tabella in documento.tables:
        prima = [c.text.strip().lower() for c in tabella.rows[0].cells]
        if all(any(attesa in cella for cella in prima) for attesa in attese):
            return tabella
    return None


def _tabella_con_etichetta(documento, etichetta: str):
    for tabella in documento.tables:
        for riga in tabella.rows:
            if riga.cells and etichetta in riga.cells[0].text.strip().lower():
                return tabella
    return None


def _patch_offerta_economica(documento) -> bool:
    """Trasforma la tabella dell'offerta economica in due cicli: prodotti e servizi.

    Il modello AD ha due blocchi identici (materiali e servizi), ciascuno con una
    riga di dettaglio da ripetere: quella riga diventa il corpo del ciclo,
    racchiusa fra i marcatori ``{%tr for ... %}`` e ``{%tr endfor %}``.
    """
    tabella = _trova_tabella(documento, ["descrizione", "quantità", "prezzo"])
    if tabella is None:
        return False

    dettagli = [
        indice for indice, riga in enumerate(tabella.rows)
        if "descrizione di dettaglio" in riga.cells[0].text.strip().lower()
    ]
    if not dettagli:
        return False

    variabili = ["prodotti", "servizi"]
    # Si parte dall'ultimo blocco: inserendo righe gli indici successivi slittano.
    for posizione, indice in reversed(list(enumerate(dettagli[:2]))):
        variabile = variabili[posizione]
        riga = tabella.rows[indice]
        _scrivi_riga(riga, ["{{ riga.descrizione_completa }}", "{{ riga.quantita }}", "{{ riga.totale }}"])
        _clona(tabella, riga, ["{%tr endfor %}"], prima=False)
        _clona(tabella, riga, ["{%tr for riga in " + variabile + " %}"], prima=True)

    def scrivi_dove(etichetta: str, testi: List[str]) -> None:
        for riga in tabella.rows:
            if etichetta in riga.cells[0].text.strip().lower():
                _scrivi_riga(riga, testi)
                return

    scrivi_dove("[descrizione generica prodotti]", ["{{ oggetto }}"])
    scrivi_dove("[descrizione generica servizi]", ["{{ nota_servizi }}"])
    scrivi_dove("totale materiali", ["TOTALE MATERIALI", "{{ totale_prodotti }}"])
    scrivi_dove("totale servizi", ["TOTALE SERVIZI", "{{ totale_servizi }}"])
    scrivi_dove("netto a voi riservato", ["Netto a Voi Riservato", "{{ totale_imponibile }}"])
    return True


def _patch_condizioni(documento) -> bool:
    tabella = _tabella_con_etichetta(documento, "tipologia di pagamento")
    if tabella is None:
        return False
    for riga in tabella.rows:
        etichetta = riga.cells[0].text.strip().lower()
        for chiave, rimpiazzo in CONDIZIONI.items():
            if etichetta.startswith(chiave) and len(riga.cells) > 1:
                _scrivi_cella(riga.cells[1], rimpiazzo)
                break
    return True


def _paragrafo_clone(paragrafo, testo: str, prima: bool):
    """Duplica un paragrafo mantenendone lo stile e ci scrive dentro."""
    from docx.text.paragraph import Paragraph

    nuovo_p = copy.deepcopy(paragrafo._p)
    if prima:
        paragrafo._p.addprevious(nuovo_p)
    else:
        paragrafo._p.addnext(nuovo_p)
    clone = Paragraph(nuovo_p, paragrafo._parent)
    _scrivi(clone, testo)
    return clone


def _ciclo_elenco(paragrafo, variabile: str, elemento: str) -> None:
    """Trasforma un punto elenco nel ciclo sui valori dell'offerta.

    I marcatori ``{%p ... %}`` cancellano il paragrafo che li contiene: apertura
    e chiusura del ciclo vanno quindi in due paragrafi propri, con in mezzo la
    voce da ripetere.
    """
    _paragrafo_clone(paragrafo, "{%p for " + elemento + " in " + variabile + " %}", prima=True)
    _scrivi(paragrafo, "{{ " + elemento + " }}")
    _paragrafo_clone(paragrafo, "{%p endfor %}", prima=False)


def _patch_paragrafi(documento) -> None:
    paragrafi = list(documento.paragraphs)
    for indice, paragrafo in enumerate(paragrafi):
        testo = paragrafo.text.strip()
        stile = (paragrafo.style.name if paragrafo.style is not None else "").lower()
        sezione = _sezione_precedente(paragrafi, indice)

        if testo == "[Descrizione]":
            _scrivi(paragrafo, "{{ oggetto }}")
        elif testo == "…" and sezione.startswith("requisiti") and "esclusioni" not in sezione:
            _ciclo_elenco(paragrafo, "requisiti", "requisito")
        elif testo == "…" and sezione.startswith("esclusioni"):
            _ciclo_elenco(paragrafo, "esclusioni", "esclusione")
        elif testo == "…":
            _scrivi(paragrafo, "{{ descrizione_fornitura }}")
        elif stile.startswith("heading") and testo.lower() == "premessa":
            successivo = paragrafi[indice + 1] if indice + 1 < len(paragrafi) else None
            if successivo is not None and not successivo.text.strip():
                _scrivi(successivo, "{{ premessa }}")


def _sezione_precedente(paragrafi, indice: int) -> str:
    for precedente in reversed(paragrafi[:indice]):
        stile = (precedente.style.name if precedente.style is not None else "").lower()
        if stile.startswith("heading") and precedente.text.strip():
            return precedente.text.strip().lower()
    return ""


def prepara(sorgente: str, destinazione: str) -> str:
    from docx import Document

    os.makedirs(os.path.dirname(os.path.abspath(destinazione)) or ".", exist_ok=True)
    _patch_copertina(sorgente, destinazione)

    documento = Document(destinazione)
    economica = _patch_offerta_economica(documento)
    condizioni = _patch_condizioni(documento)
    _patch_paragrafi(documento)
    documento.save(destinazione)

    print(f"Template preparato: {destinazione}")
    print(f"  offerta economica: {'ok' if economica else 'tabella non trovata'}")
    print(f"  condizioni di vendita: {'ok' if condizioni else 'tabella non trovata'}")
    return destinazione


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(__doc__)
        raise SystemExit(2)
    prepara(sys.argv[1], sys.argv[2])
